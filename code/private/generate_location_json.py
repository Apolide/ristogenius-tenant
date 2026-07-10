import json
import re
import argparse
from pathlib import Path
from openpyxl import load_workbook

# # LAUNCH COMMAND
# # python3 private/generate_location_json.py private/Elenco-comuni-italiani-ISTAT-2026.xlsx

# Colonne XLSX ISTAT indicate da te:
# G = Nome Comune
# K = Nome Regione
# L = Nome Provincia
COL_COMUNE = 7
COL_REGIONE = 11
COL_PROVINCIA = 12


def clean_value(value):
    """
    Pulisce i valori letti da Excel.
    """
    if value is None:
        return None

    value = str(value).strip()

    if not value:
        return None

    # Normalizza spazi multipli
    value = re.sub(r"\s+", " ", value)

    return value


def make_key(*parts):
    """
    Crea una chiave tecnica stabile per collegare province e comuni.
    Non è una vera sigla automobilistica, ma serve al Seeder per fare la mappa.
    """
    raw = "::".join(parts).lower().strip()

    replacements = {
        "à": "a",
        "è": "e",
        "é": "e",
        "ì": "i",
        "ò": "o",
        "ù": "u",
    }

    for old, new in replacements.items():
        raw = raw.replace(old, new)

    raw = re.sub(r"[^a-z0-9]+", "-", raw)
    raw = raw.strip("-")

    return raw


def generate_json_from_xlsx(input_xlsx, output_dir):
    input_xlsx = Path(input_xlsx)
    output_dir = Path(output_dir)

    if not input_xlsx.exists():
        raise FileNotFoundError(f"File XLSX non trovato: {input_xlsx}")

    output_dir.mkdir(parents=True, exist_ok=True)

    workbook = load_workbook(input_xlsx, read_only=True, data_only=True)
    sheet = workbook.active

    regions = []
    provinces = []
    comuni = []

    region_name_to_id = {}
    province_key_to_data = {}

    current_region_id = 1
    current_province_id = 1

    # Salto la prima riga perché contiene le intestazioni
    for row in sheet.iter_rows(min_row=2):
        comune_name = clean_value(row[COL_COMUNE - 1].value)
        regione_name = clean_value(row[COL_REGIONE - 1].value)
        provincia_name = clean_value(row[COL_PROVINCIA - 1].value)

        if not comune_name or not regione_name or not provincia_name:
            continue

        # Regione
        if regione_name not in region_name_to_id:
            region_name_to_id[regione_name] = current_region_id

            regions.append({
                "id": current_region_id,
                "name": regione_name
            })

            current_region_id += 1

        region_id = region_name_to_id[regione_name]

        # Provincia
        province_key = make_key(regione_name, provincia_name)

        if province_key not in province_key_to_data:
            province_key_to_data[province_key] = {
                "id": current_province_id,
                "region_id": region_id,
                "name": provincia_name,
                "sigla": province_key
            }

            provinces.append(province_key_to_data[province_key])
            current_province_id += 1

        # Comune
        comuni.append({
            "name": comune_name,
            "provincia": province_key
        })

    workbook.close()

    files = {
        "regioni_filtered.json": regions,
        "province_filtered.json": provinces,
        "comuni_filtered.json": comuni,
    }

    for filename, data in files.items():
        path = output_dir / filename

        with open(path, "w", encoding="utf-8") as f:
            json.dump(data, f, ensure_ascii=False, indent=2)

        print(f"Creato: {path} - record: {len(data)}")

    print("\nOperazione completata.")


if __name__ == "__main__":
    parser = argparse.ArgumentParser(
        description="Genera JSON regioni, province e comuni da file XLSX ISTAT."
    )

    parser.add_argument(
        "input_xlsx",
        help="Percorso del file XLSX ISTAT dei comuni italiani"
    )

    parser.add_argument(
        "--out",
        default=Path(__file__).resolve().parent / "json",
        help="Cartella di output dei JSON"
    )

    args = parser.parse_args()

    generate_json_from_xlsx(args.input_xlsx, args.out)
